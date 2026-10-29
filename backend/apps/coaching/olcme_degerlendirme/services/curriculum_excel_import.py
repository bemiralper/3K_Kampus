"""Kazanımlar sayfasının indirdiği Excel'i geri yükler.

Sütunlar: Ders kodu (isteğe bağlı), Ders, Program, Konu kodu, Konu,
Kazanım kodu, Kazanım, Alt kod, Alt metin.

Yalnızca istenen programın, dosyada adı geçen dersleri değişir.
Diğer program ve dosyada olmayan dersler durur.
"""
from __future__ import annotations

from pathlib import Path

from django.db import transaction

from apps.coaching.olcme_degerlendirme.models.curriculum import (
    Outcome,
    Subject,
    SubOutcome,
    Topic,
)

_HEADER_FIELDS = {
    'ders kodu': 'subject_code',
    'ders': 'ders',
    'program': 'program',
    'konu kodu': 'topic_code',
    'konu': 'topic_name',
    'kazanım kodu': 'outcome_code',
    'kazanim kodu': 'outcome_code',
    'kazanım': 'outcome_text',
    'kazanim': 'outcome_text',
    'alt kod': 'sub_code',
    'alt metin': 'sub_text',
}

_PROGRAMS = {
    'maarif': Topic.Program.MAARIF,
    'maarif modeli': Topic.Program.MAARIF,
    'program_2018': Topic.Program.PROGRAM_2018,
    '2018 programı': Topic.Program.PROGRAM_2018,
    '2018 programi': Topic.Program.PROGRAM_2018,
}


def _fold(value) -> str:
    text = '' if value is None else str(value)
    return text.strip().casefold().replace('i̇', 'i')


def _cell(value) -> str:
    return '' if value is None else str(value).strip()


def outcomes_header_map(header) -> dict[str, int] | None:
    """İndirilen kazanım tablosuysa sütun indeksleri. Değilse None."""
    if not header:
        return None
    found: dict[str, int] = {}
    for index, cell in enumerate(header):
        key = _HEADER_FIELDS.get(_fold(cell))
        if key and key not in found:
            found[key] = index
    required = {'ders', 'program', 'topic_code', 'topic_name', 'outcome_code', 'outcome_text'}
    if not required <= set(found):
        return None
    return found


def _program_value(raw: str) -> str:
    program = _PROGRAMS.get(_fold(raw))
    if not program:
        raise ValueError(f'Program tanınmadı: {raw or "—"}. Maarif Modeli veya 2018 Programı olmalı.')
    return program


def _resolve_subject(ders: str, subject_code: str, subjects: list[Subject]) -> Subject:
    if subject_code:
        found = next((subject for subject in subjects if subject.code.casefold() == subject_code.casefold()), None)
        if found is not None:
            return found
        raise ValueError(f'Ders kodu bulunamadı: {subject_code}')
    folded = _fold(ders)
    matches = [
        subject for subject in subjects
        if _fold(subject.display_name) == folded or _fold(subject.name) == folded
    ]
    if len(matches) == 1:
        return matches[0]
    plain = [subject for subject in matches if not subject.code.endswith(('_TYT', '_AYT', '_LGS'))]
    if len(plain) == 1:
        return plain[0]
    if matches:
        raise ValueError(f'"{ders}" birden fazla derse uyuyor. Dosyaya Ders kodu sütunu ekleyin.')
    raise ValueError(f'Ders bulunamadı: {ders}')


def parse_outcomes_workbook(
    path: str | Path,
    *,
    program: str = '',
    subject_codes: list[str] | None = None,
) -> dict:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        columns = outcomes_header_map(header)
        if columns is None:
            raise ValueError(
                'Bu dosya kazanım tablosu değil. Kazanımlar sayfasından indirilen Excel yüklenmeli.'
            )
        grouped: dict[tuple[int, str], dict] = {}
        subjects = list(Subject.objects.all())
        wanted = {code.casefold() for code in subject_codes or [] if code.strip()}
        seen_program = False
        skipped_subject = False
        for line_no, raw in enumerate(rows, start=2):
            if not raw or all(_cell(cell) == '' for cell in raw):
                continue

            def take(field: str) -> str:
                index = columns.get(field)
                if index is None or index >= len(raw):
                    return ''
                return _cell(raw[index])

            row_program = _program_value(take('program'))
            if program and row_program != program:
                continue
            ders = take('ders')
            if not ders and not take('subject_code'):
                raise ValueError(f'{line_no}. satırda ders yok.')
            subject = _resolve_subject(ders, take('subject_code'), subjects)
            if wanted and subject.code.casefold() not in wanted:
                skipped_subject = True
                continue
            seen_program = True
            topic_name = take('topic_name')
            if not topic_name:
                raise ValueError(f'{line_no}. satırda konu adı yok.')
            bucket = grouped.setdefault((subject.id, row_program), {'subject': subject, 'topics': []})
            topic_code = take('topic_code')[:30]
            topic = next(
                (item for item in bucket['topics'] if item['code'] == topic_code and item['name'] == topic_name[:200]),
                None,
            )
            if topic is None:
                topic = {'code': topic_code, 'name': topic_name[:200], 'outcomes': []}
                bucket['topics'].append(topic)
            outcome_code = take('outcome_code')[:50]
            outcome_text = take('outcome_text')
            if not outcome_code and not outcome_text:
                continue
            outcome = next(
                (
                    item for item in topic['outcomes']
                    if item['code'] == outcome_code and item['text'] == outcome_text
                ),
                None,
            )
            if outcome is None:
                outcome = {'code': outcome_code, 'text': outcome_text, 'sub_outcomes': []}
                topic['outcomes'].append(outcome)
            sub_code = take('sub_code')[:50]
            sub_text = take('sub_text')
            if sub_code or sub_text:
                outcome['sub_outcomes'].append({'code': sub_code, 'text': sub_text})
        if program and not seen_program:
            if skipped_subject:
                raise ValueError('Seçilen dersler bu dosyada yok.')
            label = 'Maarif Modeli' if program == Topic.Program.MAARIF else '2018 Programı'
            raise ValueError(f'Bu dosyada {label} kazanımı yok.')
        return {
            key: value['topics']
            for key, value in grouped.items()
        } | {'_subjects': {key: value['subject'] for key, value in grouped.items()}}
    finally:
        wb.close()


def _protected_topic_ids(topic_ids: list[int]) -> set[int]:
    from apps.mobile_identity.models import IslenenKonu

    if not topic_ids:
        return set()
    return set(
        IslenenKonu.objects.filter(topic_id__in=topic_ids).values_list('topic_id', flat=True)
    )


def _write_outcomes(topic: Topic, outcomes: list[dict]) -> None:
    from apps.mobile_identity.models import IslenenKazanim

    protected = {
        row.code: row
        for row in Outcome.objects.filter(
            topic=topic,
            id__in=IslenenKazanim.objects.filter(outcome__topic=topic).values('outcome_id'),
        )
    }
    topic.outcomes.exclude(id__in=[row.id for row in protected.values()]).delete()
    for index, outcome_data in enumerate(outcomes):
        code = outcome_data['code']
        existing = protected.get(code)
        if existing is None:
            existing = Outcome.objects.create(
                topic=topic, code=code, text=outcome_data['text'], order=index,
            )
        else:
            existing.text = outcome_data['text']
            existing.order = index
            existing.save(update_fields=['text', 'order'])
        existing.sub_outcomes.all().delete()
        if outcome_data['sub_outcomes']:
            SubOutcome.objects.bulk_create([
                SubOutcome(
                    outcome=existing, code=sub['code'], text=sub['text'], order=sub_index,
                )
                for sub_index, sub in enumerate(outcome_data['sub_outcomes'])
            ])


def import_outcomes_workbook(
    path: str | Path,
    *,
    program: str = '',
    subject_codes: list[str] | None = None,
) -> dict:
    parsed = parse_outcomes_workbook(path, program=program, subject_codes=subject_codes)
    subjects = parsed.pop('_subjects')
    stats = {
        'subjects': len(subjects),
        'topics': 0,
        'outcomes': 0,
        'sub_outcomes': 0,
        'kept_topics': 0,
    }
    with transaction.atomic():
        for key, topics in parsed.items():
            subject_id, row_program = key
            subject = subjects[(subject_id, row_program)]
            existing = Topic.objects.filter(subject=subject, program=row_program)
            kept = {
                (topic.code, topic.name): topic
                for topic in existing.filter(id__in=_protected_topic_ids(
                    list(existing.values_list('id', flat=True)),
                ))
            }
            stats['kept_topics'] += len(kept)
            existing.exclude(id__in=[topic.id for topic in kept.values()]).delete()
            order_base = (
                Topic.objects.filter(subject=subject).order_by('-order')
                .values_list('order', flat=True).first() or 0
            )
            for index, topic_data in enumerate(topics, start=1):
                topic = kept.get((topic_data['code'], topic_data['name']))
                if topic is None:
                    topic = Topic.objects.create(
                        subject=subject,
                        program=row_program,
                        code=topic_data['code'],
                        name=topic_data['name'],
                        order=order_base + index,
                    )
                else:
                    topic.order = order_base + index
                    topic.save(update_fields=['order'])
                _write_outcomes(topic, topic_data['outcomes'])
                stats['topics'] += 1
                stats['outcomes'] += len(topic_data['outcomes'])
                stats['sub_outcomes'] += sum(
                    len(outcome['sub_outcomes']) for outcome in topic_data['outcomes']
                )
    return stats
