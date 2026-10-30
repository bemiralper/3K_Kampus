"""Günü birlik deneme kaydı: kısa kayıt, TC eşleme, kitleye sızmama."""
import json
from datetime import timedelta
from io import BytesIO
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.test import TestCase
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from rest_framework.test import APIClient

from apps.coaching.olcme_degerlendirme.models import Exam, ExamParticipant
from apps.coaching.olcme_degerlendirme.services.exam_roster import (
    replace_auto_participants,
    resolve_exam_candidates,
)
from apps.coaching.olcme_degerlendirme.views.analysis_views import (
    _gunu_birlik_labels,
    _program_label,
)
from apps.coaching.tests.olcme_helpers import grant_olcme_write
from apps.egitim_paketleri.models import EkHizmet
from apps.egitim_tanimlari.models import SinifSeviyesi
from apps.egitim_yili.domain.models import EgitimYili
from apps.kurum.domain.models import Kurum
from apps.ogrenci.domain.models import Ogrenci, OgrenciAdres, OgrenciKayit, OgrenciVeli
from apps.ogrenci.interfaces.list_helpers import build_kayit_queryset
from apps.ogrenci.services.kayit_turu import GUNU_BIRLIK_CODE
from apps.ogrenci_kayit.domain.models import LookupCategory, LookupOption
from apps.roller.models import Role, UserRole
from apps.roller.seed import ensure_default_roles
from apps.sube.domain.models import Sube

User = get_user_model()
URL = '/api/ogrenci-kayit/gunu-birlik/'


class GunuBirlikApiTests(TestCase):
    def setUp(self):
        self.kurum = Kurum.objects.create(ad='Günü Kurum', kod='GUNU')
        self.sube = Sube.objects.create(kurum=self.kurum, ad='Merkez', kod='GUNU-M')
        self.yil = EgitimYili.objects.create(baslangic_yil=2025, bitis_yil=2026, aktif_mi=True)
        self.sev12 = SinifSeviyesi.objects.create(
            kurum=self.kurum, sube=self.sube, ad='12. Sınıf', kod='12',
        )
        self.sev8 = SinifSeviyesi.objects.create(
            kurum=self.kurum, sube=self.sube, ad='8. Sınıf', kod='8',
        )
        today = timezone.localdate()
        self.exam = Exam.objects.create(
            name='TYT Deneme', exam_type='YKS_TYT', exam_date=today + timedelta(days=20),
            kurum=self.kurum, sube=self.sube, egitim_yili=self.yil,
        )
        self.exam_b = Exam.objects.create(
            name='Başka Deneme', exam_type='YKS_TYT', exam_date=today + timedelta(days=40),
            kurum=self.kurum, sube=self.sube, egitim_yili=self.yil,
        )
        self.user = User.objects.create_user(username='gunu_memur', password='x')
        grant_olcme_write(self.user, self.kurum)
        self.client = APIClient()
        self.client.force_authenticate(user=self.user)
        self.headers = {
            'HTTP_X_KURUM_ID': str(self.kurum.id),
            'HTTP_X_SUBE_ID': str(self.sube.id),
            'HTTP_X_EGITIMYILI_ID': str(self.yil.id),
        }

    def _row(self, ad='Ayşe', soyad='Yılmaz', tc='10000000001', seviye=None, telefon='05321112233'):
        return {
            'ad': ad,
            'soyad': soyad,
            'tc_kimlik_no': tc,
            'telefon': telefon,
            'sinif_seviyesi_id': (seviye or self.sev12).id,
        }

    def _post(self, rows, *, exam=None, commit=True):
        return self.client.post(
            URL,
            {
                'exam_id': (exam or self.exam).id,
                'commit': commit,
                'ogrenciler': rows,
            },
            format='json',
            **self.headers,
        )

    def test_yeni_tc_yillik_kayit_acmaz_ve_secilen_sinava_duser(self):
        res = self._post([self._row()])
        self.assertEqual(res.status_code, 200, res.content)
        body = res.json()
        self.assertEqual(body['rows'][0]['status'], 'yeni')
        ogrenci = Ogrenci.objects.get(tc_kimlik_no='10000000001')
        self.assertEqual(ogrenci.kayit_turu, GUNU_BIRLIK_CODE)
        self.assertFalse(OgrenciKayit.objects.filter(ogrenci=ogrenci).exists())
        part = ExamParticipant.objects.get(exam=self.exam, student=ogrenci)
        self.assertEqual(part.source, ExamParticipant.Source.MANUAL)
        self.assertEqual(part.sinif_seviyesi_id, self.sev12.id)
        self.assertFalse(ExamParticipant.objects.filter(exam=self.exam_b, student=ogrenci).exists())
        label = _gunu_birlik_labels(self.exam, [ogrenci.id])[ogrenci.id]
        self.assertEqual(label, '12. Sınıf · Tek deneme')
        self.assertEqual(_program_label('', False, False, label), label)

    def test_ayni_tc_ikinci_sinavda_ogrenci_cogalmaz(self):
        self._post([self._row()])
        res = self._post([self._row(ad='Ayse', soyad='Yilmaz')], exam=self.exam_b)
        self.assertEqual(res.status_code, 200, res.content)
        self.assertEqual(res.json()['rows'][0]['status'], 'mevcut_misafir')
        self.assertEqual(Ogrenci.objects.filter(tc_kimlik_no='10000000001').count(), 1)
        ogrenci = Ogrenci.objects.get(tc_kimlik_no='10000000001')
        self.assertEqual(ogrenci.kayit_turu, GUNU_BIRLIK_CODE)
        self.assertTrue(ExamParticipant.objects.filter(exam=self.exam, student=ogrenci).exists())
        self.assertTrue(ExamParticipant.objects.filter(exam=self.exam_b, student=ogrenci).exists())

    def test_asil_ogrenci_turu_bozulmaz_sinava_eklenir(self):
        asil = Ogrenci.objects.create(
            kurum=self.kurum, sube=self.sube, ad='Ali', soyad='Kaya',
            tc_kimlik_no='10000000002', kayit_turu='asil', aktif_mi=True,
        )
        OgrenciKayit.objects.create(
            ogrenci=asil, egitim_yili=self.yil, kurum=self.kurum, sube=self.sube,
            sinif_seviyesi=self.sev12, aktif_mi=True,
        )
        res = self._post([self._row(ad='Ali', soyad='Kaya', tc='10000000002')])
        self.assertEqual(res.status_code, 200, res.content)
        self.assertEqual(res.json()['rows'][0]['status'], 'kayitli_ogrenci')
        self.assertEqual(res.json()['rows'][0]['message'], 'Kayıtlı öğrenci, sınava eklendi.')
        asil.refresh_from_db()
        self.assertEqual(asil.kayit_turu, 'asil')
        self.assertEqual(Ogrenci.objects.filter(tc_kimlik_no='10000000002').count(), 1)
        self.assertTrue(ExamParticipant.objects.filter(exam=self.exam, student=asil).exists())

    def test_seviye_kitlesi_gunu_birlik_kaydi_almaz(self):
        self._post([self._row()])
        guest = Ogrenci.objects.get(tc_kimlik_no='10000000001')
        OgrenciKayit.objects.create(
            ogrenci=guest, egitim_yili=self.yil, kurum=self.kurum, sube=self.sube,
            sinif_seviyesi=self.sev12, aktif_mi=True,
        )
        classmate = Ogrenci.objects.create(
            kurum=self.kurum, sube=self.sube, ad='Sinif', soyad='Ogrenci', kayit_turu='asil',
        )
        OgrenciKayit.objects.create(
            ogrenci=classmate, egitim_yili=self.yil, kurum=self.kurum, sube=self.sube,
            sinif_seviyesi=self.sev12, aktif_mi=True,
        )
        recs = resolve_exam_candidates(
            kurum_id=self.kurum.id, sube_id=self.sube.id, egitim_yili_id=self.yil.id,
            seviye_ids=[self.sev12.id],
        )
        ids = {r.student_id for r in recs}
        self.assertIn(classmate.id, ids)
        self.assertNotIn(guest.id, ids)
        replace_auto_participants(self.exam_b, recs)
        self.assertFalse(ExamParticipant.objects.filter(exam=self.exam_b, student=guest).exists())
        self.assertTrue(ExamParticipant.objects.filter(exam=self.exam, student=guest).exists())

    def test_bozuk_tc_kaydi_durdurur_saglam_satir_ayrilir(self):
        rows = [
            self._row(tc='123'),
            self._row(ad='Veli', soyad='Demir', tc='10000000003'),
        ]
        preview = self._post(rows, commit=False)
        self.assertEqual(preview.status_code, 200, preview.content)
        statuses = [row['status'] for row in preview.json()['rows']]
        self.assertEqual(statuses, ['hata', 'yeni'])
        self.assertFalse(Ogrenci.objects.filter(tc_kimlik_no='10000000003').exists())

        saved = self._post(rows, commit=True)
        self.assertEqual(saved.status_code, 200, saved.content)
        saved_rows = saved.json()['rows']
        self.assertEqual(saved_rows[0]['status'], 'hata')
        self.assertEqual(saved_rows[1]['status'], 'yeni')
        self.assertTrue(Ogrenci.objects.filter(tc_kimlik_no='10000000003', kayit_turu=GUNU_BIRLIK_CODE).exists())
        self.assertFalse(Ogrenci.objects.filter(tc_kimlik_no='123').exists())

    def test_excel_onizleme_bozuk_satiri_ayirir(self):
        wb = Workbook()
        ws = wb.active
        ws.append(['Ad', 'Soyad', 'TC', 'Telefon', 'Seviye'])
        ws.append(['Kötü', 'Satır', '123', '', '12'])
        ws.append(['Veli', 'Demir', '10000000004', '05320000000', '12'])
        buf = BytesIO()
        wb.save(buf)
        buf.seek(0)
        buf.name = 'liste.xlsx'
        res = self.client.post(
            f'{URL}excel/',
            {'exam_id': self.exam.id, 'file': buf},
            format='multipart',
            **self.headers,
        )
        self.assertEqual(res.status_code, 200, res.content)
        self.assertEqual([row['status'] for row in res.json()['rows']], ['hata', 'yeni'])
        self.assertFalse(Ogrenci.objects.filter(tc_kimlik_no='10000000004').exists())
        self.assertFalse(res.json()['committed'])

    def test_sablon_xlsx(self):
        res = self.client.get(f'{URL}sablon/', **self.headers)
        self.assertEqual(res.status_code, 200)
        wb = load_workbook(BytesIO(res.content))
        self.assertEqual([c.value for c in next(wb.active.iter_rows(max_row=1))], ['Ad', 'Soyad', 'TC', 'Telefon', 'Seviye'])

    def test_secenekler_gecmis_ve_baska_seviyeyi_gizler(self):
        today = timezone.localdate()
        self.exam.exam_date = today - timedelta(days=1)
        self.exam.save(update_fields=['exam_date'])
        lgs = Exam.objects.create(
            name='LGS Deneme', exam_type='LGS', exam_date=today + timedelta(days=5),
            kurum=self.kurum, sube=self.sube, egitim_yili=self.yil,
        )
        res = self.client.get(f'{URL}secenekler/?sinif_seviyesi={self.sev12.id}', **self.headers)
        self.assertEqual(res.status_code, 200, res.content)
        ids = [item['id'] for item in res.json()['sinavlar']]
        self.assertIn(self.exam_b.id, ids)
        self.assertNotIn(self.exam.id, ids)
        self.assertNotIn(lgs.id, ids)

        res8 = self.client.get(f'{URL}secenekler/?sinif_seviyesi={self.sev8.id}', **self.headers)
        ids8 = [item['id'] for item in res8.json()['sinavlar']]
        self.assertEqual(ids8, [lgs.id])

    def test_ilk_oturum_gecse_sonraki_oturum_listelenir(self):
        from apps.coaching.olcme_degerlendirme.models import ExamSessionModel

        today = timezone.localdate()
        ExamSessionModel.objects.create(
            exam=self.exam, name='1. Oturum', order=1,
            session_date=today - timedelta(days=2),
        )
        later = ExamSessionModel.objects.create(
            exam=self.exam, name='2. Oturum', order=2,
            session_date=today + timedelta(days=5),
        )
        self.exam.exam_date = today - timedelta(days=2)
        self.exam.save(update_fields=['exam_date'])
        res = self.client.get(f'{URL}secenekler/?sinif_seviyesi={self.sev12.id}', **self.headers)
        self.assertEqual(res.status_code, 200, res.content)
        item = next(row for row in res.json()['sinavlar'] if row['id'] == self.exam.id)
        self.assertEqual([session['id'] for session in item['sessions']], [later.id])
        self.assertEqual(item['exam_date'], (today + timedelta(days=5)).isoformat())

    def test_kitledeki_seviye_sinav_turunu_ezer(self):
        from apps.coaching.olcme_degerlendirme.models import ExamAudience

        ExamAudience.objects.create(exam=self.exam_b, sinif_seviyesi=self.sev8)
        res12 = self.client.get(f'{URL}secenekler/?sinif_seviyesi={self.sev12.id}', **self.headers)
        ids12 = [item['id'] for item in res12.json()['sinavlar']]
        self.assertNotIn(self.exam_b.id, ids12)
        self.assertIn(self.exam.id, ids12)
        res8 = self.client.get(f'{URL}secenekler/?sinif_seviyesi={self.sev8.id}', **self.headers)
        ids8 = [item['id'] for item in res8.json()['sinavlar']]
        self.assertEqual(ids8, [self.exam_b.id])

    def test_kisa_kayit_uyumsuz_seviyeyi_yazmaz(self):
        res = self._post([self._row(seviye=self.sev8, tc='10000000008')])
        self.assertEqual(res.status_code, 200, res.content)
        self.assertEqual(res.json()['rows'][0]['status'], 'hata')
        self.assertIn('sınıf', res.json()['rows'][0]['message'])
        self.assertFalse(Ogrenci.objects.filter(tc_kimlik_no='10000000008').exists())

    def test_gecmis_sinav_kisa_kaydi_reddeder(self):
        self.exam.exam_date = timezone.localdate() - timedelta(days=2)
        self.exam.save(update_fields=['exam_date'])
        res = self._post([self._row()])
        self.assertEqual(res.status_code, 400, res.content)
        self.assertIn('geçmiş', res.json()['error'])


class GunuBirlikTamKayitTests(TestCase):
    def setUp(self):
        ensure_default_roles()
        self.kurum = Kurum.objects.create(ad='Yükselt Kurum', kod='YUK')
        self.sube = Sube.objects.create(kurum=self.kurum, ad='Merkez', kod='YUK-M')
        self.yil = EgitimYili.objects.create(baslangic_yil=2025, bitis_yil=2026, aktif_mi=True)
        self.seviye = SinifSeviyesi.objects.create(
            kurum=self.kurum, sube=self.sube, ad='8. Sınıf', kod='8',
        )
        self.hizmet = EkHizmet.objects.create(
            ad='Kütüphane', kod='KUT', hizmet_turu='kutuphane',
            kurum=self.kurum, sube=self.sube, egitim_yili=self.yil,
            brut_fiyat=1000, kdv_orani=10,
        )
        self.user = User.objects.create_user(username='yuk_memur', password='x')
        UserRole.objects.create(
            user=self.user,
            role=Role.objects.get(code='kurum_yoneticisi'),
            kurum=self.kurum,
            must_change_password=False,
        )
        self.client = APIClient()
        self.client.force_authenticate(user=self.user)
        self.headers = {
            'HTTP_X_KURUM_ID': str(self.kurum.id),
            'HTTP_X_SUBE_ID': str(self.sube.id),
            'HTTP_X_EGITIMYILI_ID': str(self.yil.id),
        }
        self.guest = Ogrenci.objects.create(
            kurum=self.kurum, sube=self.sube, ad='Eski', soyad='Misafir',
            tc_kimlik_no='10000000009', telefon='05329998877',
            kayit_turu=GUNU_BIRLIK_CODE, aktif_mi=True,
        )

    def test_tam_kayit_gunu_birlik_turunu_gunceller(self):
        category, _ = LookupCategory.objects.get_or_create(
            code='registration_type', defaults={'name': 'Kayıt Türü'},
        )
        asil, _ = LookupOption.objects.get_or_create(
            category=category, code='asil', defaults={'label': 'Asil', 'order': 1},
        )
        res = self.client.post(
            '/api/ogrenci-kayit/register/',
            data=json.dumps({
                'student': {
                    'tc_kimlik_no': '10000000009',
                    'ad': 'Yeni',
                    'soyad': 'Öğrenci',
                    'dogum_tarihi': '2012-03-01',
                    'kayit_turu': asil.id,
                },
                'enrollment': {
                    'egitim_yili': self.yil.id,
                    'sinif_seviyesi': self.seviye.id,
                },
                'address': {},
                'guardians': [],
                'package': {'ek_hizmet_ids': [self.hizmet.id]},
            }),
            content_type='application/json',
            **self.headers,
        )
        self.assertEqual(res.status_code, 201, res.content)
        self.guest.refresh_from_db()
        self.assertEqual(self.guest.kayit_turu, 'asil')
        self.assertEqual(self.guest.ad, 'Yeni')
        self.assertTrue(OgrenciKayit.objects.filter(ogrenci=self.guest, egitim_yili=self.yil).exists())
        self.assertEqual(Ogrenci.objects.filter(tc_kimlik_no='10000000009').count(), 1)

    def _register(self, payload):
        return self.client.post(
            '/api/ogrenci-kayit/register/',
            data=json.dumps(payload),
            content_type='application/json',
            **self.headers,
        )

    def _walk_payload(self, **student):
        base = {
            'tc_kimlik_no': '10000000046',
            'ad': 'Deneme',
            'soyad': 'Misafir',
            'dogum_tarihi': '2011-04-04',
            'telefon': '05321110046',
        }
        base.update(student)
        return {
            'student': base,
            'enrollment': {
                'egitim_yili': self.yil.id,
                'sinif_seviyesi': self.seviye.id,
                'ogrenci_no': 'GB-46',
            },
            'address': {'acik_adres': 'Deneme sokak 1'},
            'guardians': [{
                'ad': 'Veli',
                'soyad': 'Misafir',
                'tc_kimlik_no': '10000000047',
                'telefon': '05321110047',
            }],
            'veliSecimi': 'add',
            'package': {'ek_hizmet_ids': [self.hizmet.id]},
            'gunu_birlik': {'exam_id': self.exam.id},
        }

    def _nakit(self):
        from apps.finans.constants.payment_types import OdemeYontemiTipi
        from apps.finans.domain.financial_account import MaliHesap
        from apps.finans.domain.payment_method import OdemeYontemi

        hesap = MaliHesap.objects.create(sube=self.sube, ad='Nakit Kasa', tip='kasa')
        yontem = OdemeYontemi.objects.create(
            mali_hesap=hesap, kurum=self.kurum, ad='Nakit', tip=OdemeYontemiTipi.NAKIT,
        )
        return hesap, yontem

    def test_wizard_veli_okul_ve_secilen_sinavi_yazar_listeye_dusmez(self):
        from apps.coaching.olcme_degerlendirme.models import Exam, ExamParticipant

        self.exam = Exam.objects.create(
            name='LGS Deneme', exam_type='LGS', exam_date=timezone.localdate() + timedelta(days=20),
            kurum=self.kurum, sube=self.sube, egitim_yili=self.yil,
        )
        other = Exam.objects.create(
            name='Başka Deneme', exam_type='LGS', exam_date=timezone.localdate() + timedelta(days=40),
            kurum=self.kurum, sube=self.sube, egitim_yili=self.yil,
        )
        res = self._register(self._walk_payload())
        self.assertEqual(res.status_code, 201, res.content)
        self.assertIn('sınava eklendi', res.json()['message'])
        ogrenci = Ogrenci.objects.get(tc_kimlik_no='10000000046')
        self.assertEqual(ogrenci.kayit_turu, GUNU_BIRLIK_CODE)
        self.assertTrue(OgrenciKayit.objects.filter(ogrenci=ogrenci, egitim_yili=self.yil).exists())
        self.assertTrue(OgrenciAdres.objects.filter(ogrenci=ogrenci, adres='Deneme sokak 1').exists())
        self.assertTrue(OgrenciVeli.objects.filter(ogrenci=ogrenci, ad='Veli', telefon='05321110047').exists())
        self.assertTrue(ExamParticipant.objects.filter(exam=self.exam, student=ogrenci).exists())
        self.assertFalse(ExamParticipant.objects.filter(exam=other, student=ogrenci).exists())
        ctx = {
            'kurum_id': self.kurum.id,
            'sube_id': self.sube.id,
            'egitim_yili_id': self.yil.id,
        }
        params = {
            'q': '', 'all_years': False, 'durum': 'aktif',
            'sinif_seviyesi_ids': [], 'giris_turu': None, 'kayit_turu': None,
            'cinsiyet': None, 'paket_id': None, 'paket_turu': None, 'kalemler': [],
            'sinif_ids': [], 'school_ids': [], 'alan_ids': [], 'coach_ids': [],
            'kayit_tarihi_bas': None, 'kayit_tarihi_bit': None,
            'sort': 'created_at_desc', 'page': 1, 'page_size': 25,
        }
        qs, _ = build_kayit_queryset(ctx, params)
        self.assertNotIn(ogrenci.id, set(qs.values_list('ogrenci_id', flat=True)))
        params['kayit_turu'] = GUNU_BIRLIK_CODE
        listed, _ = build_kayit_queryset(ctx, params)
        self.assertIn(ogrenci.id, set(listed.values_list('ogrenci_id', flat=True)))

    def test_ana_liste_aramasi_deneme_ogrencisini_ayirir(self):
        self.exam = Exam.objects.create(
            name='Liste Deneme', exam_type='LGS',
            exam_date=timezone.localdate() + timedelta(days=10),
            kurum=self.kurum, sube=self.sube, egitim_yili=self.yil,
        )
        payload = self._walk_payload()
        payload['gunu_birlik'] = {'kayit': True}
        payload['package'] = {}
        res = self._register(payload)
        self.assertEqual(res.status_code, 201, res.content)
        listed = self.client.get('/ogrenciler/api/list/?q=Misafir', **self.headers)
        self.assertEqual(listed.status_code, 200, listed.content)
        body = listed.json()
        self.assertEqual(body.get('total_count'), 0)
        self.assertGreaterEqual(body.get('deneme_count') or 0, 1)

    def test_gunu_birlik_once_ogrenci_sonra_listeden_deneme(self):
        from apps.odeme_takip.domain.models import Sozlesme
        from apps.ogrenci_kayit.application.deneme_ucret import MARKER
        from apps.ogrenci_kayit.application.gunu_birlik import list_misafirler

        self.exam = Exam.objects.create(
            name='LGS Deneme', exam_type='LGS', exam_date=timezone.localdate() + timedelta(days=20),
            kurum=self.kurum, sube=self.sube, egitim_yili=self.yil,
        )
        payload = self._walk_payload()
        payload['gunu_birlik'] = {'kayit': True}
        payload['package'] = {}
        res = self._register(payload)
        self.assertEqual(res.status_code, 201, res.content)
        body = res.json()
        self.assertNotIn('exam_id', body)
        self.assertIn('açıldı', body['message'])
        ogrenci = Ogrenci.objects.get(tc_kimlik_no='10000000046')
        self.assertEqual(ogrenci.kayit_turu, GUNU_BIRLIK_CODE)
        self.assertFalse(ExamParticipant.objects.filter(student=ogrenci).exists())
        rows = list_misafirler(self.kurum.id, self.sube.id, q='Deneme')
        row = next(item for item in rows if item['id'] == ogrenci.id)
        self.assertEqual(row['sinavlar'], [])
        self.assertEqual(row['sinif_seviyesi_id'], self.seviye.id)
        self.assertEqual(row['sinif_seviyesi'], '8. Sınıf')

        _hesap, yontem = self._nakit()
        kayit = self.client.post(
            '/api/ogrenci-kayit/gunu-birlik/sinav/',
            data=json.dumps({
                'islem': 'kayit',
                'ogrenci_id': ogrenci.id,
                'exam_id': self.exam.id,
                'deneme_ucreti': '250',
                'odeme_yontemi_id': yontem.id,
                'mali_hesap_id': _hesap.id,
            }),
            content_type='application/json',
            **self.headers,
        )
        self.assertEqual(kayit.status_code, 200, kayit.content)
        self.assertTrue(ExamParticipant.objects.filter(exam=self.exam, student=ogrenci).exists())
        sozlesme = Sozlesme.objects.get(ogrenci=ogrenci, notlar__contains=MARKER)
        self.assertEqual(int(sozlesme.net_tutar), 250)

        tekrar = self.client.post(
            '/api/ogrenci-kayit/gunu-birlik/sinav/',
            data=json.dumps({
                'islem': 'kayit',
                'ogrenci_id': ogrenci.id,
                'exam_id': self.exam.id,
            }),
            content_type='application/json',
            **self.headers,
        )
        self.assertEqual(tekrar.status_code, 400, tekrar.content)
        self.assertIn('zaten', tekrar.json()['detail'])

    def test_kayitli_ogrenci_turu_degismez_sinava_eklenir(self):
        from apps.coaching.olcme_degerlendirme.models import Exam, ExamParticipant

        self.exam = Exam.objects.create(
            name='LGS Deneme', exam_type='LGS', exam_date=timezone.localdate() + timedelta(days=20),
            kurum=self.kurum, sube=self.sube, egitim_yili=self.yil,
        )
        asil = Ogrenci.objects.create(
            kurum=self.kurum, sube=self.sube, ad='Asıl', soyad='Öğrenci',
            tc_kimlik_no='10000000048', telefon='05321110048',
            kayit_turu='asil', aktif_mi=True,
        )
        OgrenciKayit.objects.create(
            ogrenci=asil, kurum=self.kurum, sube=self.sube, egitim_yili=self.yil, aktif_mi=True,
        )
        res = self._register(self._walk_payload(
            tc_kimlik_no='10000000048', ad='Asıl', soyad='Öğrenci', telefon='05321110048',
        ))
        self.assertEqual(res.status_code, 201, res.content)
        self.assertEqual(res.json()['message'], 'Kayıtlı öğrenci, sınava eklendi.')
        asil.refresh_from_db()
        self.assertEqual(asil.kayit_turu, 'asil')
        self.assertEqual(asil.ad, 'Asıl')
        self.assertEqual(OgrenciKayit.objects.filter(ogrenci=asil).count(), 1)
        self.assertTrue(ExamParticipant.objects.filter(exam=self.exam, student=asil).exists())

    def test_paketsiz_kayit_giris_belgesi_uretir(self):
        from apps.communication.application.communication_service import SendResult
        from apps.ogrenci.domain.models import OgrenciEgitimPaketi
        from apps.ogrenci_kayit.application.giris_belgesi import build_giris_belgesi_html, load_giris_context

        self.exam = Exam.objects.create(
            name='LGS Deneme', exam_type='LGS', exam_date=timezone.localdate() + timedelta(days=12),
            kurum=self.kurum, sube=self.sube, egitim_yili=self.yil,
        )
        payload = self._walk_payload()
        payload['package'] = {}
        res = self._register(payload)
        self.assertEqual(res.status_code, 201, res.content)
        body = res.json()
        self.assertEqual(body['exam_id'], self.exam.id)
        self.assertTrue(body['participant_id'])
        ogrenci = Ogrenci.objects.get(tc_kimlik_no='10000000046')
        self.assertFalse(OgrenciEgitimPaketi.objects.filter(ogrenci=ogrenci).exists())

        ctx, reason = load_giris_context(self.kurum.id, self.sube.id, self.exam.id, ogrenci.id)
        self.assertEqual(reason, '')
        html = build_giris_belgesi_html(ctx)
        self.assertIn('SINAV GİRİŞ BELGESİ', html)
        self.assertIn('Deneme Misafir', html)
        self.assertNotIn('class="student"', html)
        self.assertIn('10000000046', html)
        self.assertNotIn('Tarihi gelmemiş', html)
        self.assertNotIn('Girmemiş', html)
        self.assertIn('page-foot', html)
        self.assertIn('T.C. KİMLİK NO', html)
        self.assertIn('LGS Deneme', html)
        self.assertIn('HAZIR BULUNMA', html)
        self.assertIn('Kurşun kalem, silgi ve kalemtıraşınızı getiriniz.', html)
        self.assertIn('Telefon ve akıllı saatlerinizi görevliye teslim ediniz.', html)
        self.assertIn('Belirtilen salon ve sıra numarasına göre yerleşiniz.', html)
        self.assertIn('Sınavdan 15 dakika önce salonda hazır bulununuz.', html)
        self.assertIn('Sınav sonuç bilgilendirmesi', html)
        self.assertIn('WhatsApp numaralarına gönderilecektir', html)
        self.assertIn('class="seatbox"', html)
        self.assertEqual(body['seat_status'], 'no_room')

        pdf_bytes = b'%PDF-1.4\n' + (b'x' * 3000)
        with patch(
            'apps.communication.application.html_to_pdf.render_html_to_pdf',
            return_value=pdf_bytes,
        ):
            pdf = self.client.get(
                f'/api/ogrenci-kayit/gunu-birlik/giris-belgesi/?exam_id={self.exam.id}&ogrenci_id={ogrenci.id}',
                **self.headers,
            )
        self.assertEqual(pdf.status_code, 200, pdf.content[:200])
        self.assertEqual(pdf['Content-Type'], 'application/pdf')
        self.assertTrue(pdf.content.startswith(b'%PDF'))

        with patch(
            'apps.communication.application.notification_dispatcher.dispatch_event',
            return_value=SendResult(success=True),
        ) as notify, patch(
            'apps.communication.application.html_to_pdf.render_html_to_pdf',
            return_value=pdf_bytes,
        ):
            sent = self.client.post(
                '/api/ogrenci-kayit/gunu-birlik/giris-belgesi/whatsapp/',
                data=json.dumps({
                    'exam_id': self.exam.id,
                    'ogrenci_id': ogrenci.id,
                    'alicilar': ['ogrenci', 'veli'],
                }),
                content_type='application/json',
                **self.headers,
            )
        self.assertEqual(sent.status_code, 200, sent.content)
        self.assertEqual(sent.json()['sent'], 2)
        self.assertEqual(notify.call_count, 2)
        self.assertEqual(notify.call_args_list[0].args[1], 'sinav.giris_belgesi')

    def test_ayin_sinav_tablosu_seviyeye_ve_katilima_gore(self):
        from apps.coaching.olcme_degerlendirme.models import ExamParticipant, ExamSessionModel
        from apps.ogrenci_kayit.application.giris_belgesi import build_giris_belgesi_html, load_giris_context

        today = timezone.localdate()
        self.exam = Exam.objects.create(
            name='LGS Deneme', exam_type='LGS', exam_date=today,
            kurum=self.kurum, sube=self.sube, egitim_yili=self.yil,
        )
        res = self._register(self._walk_payload())
        self.assertEqual(res.status_code, 201, res.content)
        ogrenci = Ogrenci.objects.get(tc_kimlik_no='10000000046')

        missed = Exam.objects.create(
            name='Kaçırılan LGS', exam_type='LGS', exam_date=today.replace(day=1),
            kurum=self.kurum, sube=self.sube, egitim_yili=self.yil,
        )
        ExamSessionModel.objects.create(exam=missed, name='Sabah', order=1, session_date=today.replace(day=1))
        sat = Exam.objects.create(
            name='Girilen LGS', exam_type='LGS', exam_date=today.replace(day=1),
            kurum=self.kurum, sube=self.sube, egitim_yili=self.yil,
        )
        ExamParticipant.objects.create(exam=sat, student=ogrenci, sinif_seviyesi=self.seviye)
        Exam.objects.create(
            name='TYT Başka', exam_type='YKS_TYT', exam_date=today,
            kurum=self.kurum, sube=self.sube, egitim_yili=self.yil,
        )
        later = (today.replace(day=28) + timedelta(days=10)).replace(day=1)
        Exam.objects.create(
            name='Gelecek Ay LGS', exam_type='LGS', exam_date=later,
            kurum=self.kurum, sube=self.sube, egitim_yili=self.yil,
        )
        from calendar import monthrange
        last_day = today.replace(day=monthrange(today.year, today.month)[1])
        spill_day = last_day + timedelta(days=2)
        spanning = Exam.objects.create(
            name='Sarkan LGS', exam_type='LGS', exam_date=last_day,
            kurum=self.kurum, sube=self.sube, egitim_yili=self.yil,
        )
        ExamSessionModel.objects.create(
            exam=spanning, name='Ilk Gun', order=1, session_date=last_day,
        )
        ExamSessionModel.objects.create(
            exam=spanning, name='Sonraki Ay', order=2, session_date=spill_day,
        )

        ctx, reason = load_giris_context(self.kurum.id, self.sube.id, self.exam.id, ogrenci.id)
        self.assertEqual(reason, '')
        html = build_giris_belgesi_html(ctx)
        self.assertIn('KAÇIRILAN LGS', html)
        self.assertEqual(html.count('KAÇIRILAN LGS'), 1)
        self.assertNotIn('Sabah', html)
        self.assertIn('GİRİLEN LGS', html)
        self.assertNotIn('TYT Başka', html)
        self.assertNotIn('TYT BAŞKA', html)
        self.assertNotIn('Gelecek Ay LGS', html)
        self.assertIn('SARKAN LGS', html)
        self.assertNotIn('Sonraki Ay', html)
        from apps.ogrenci_kayit.application.giris_belgesi import _date_span
        self.assertIn(_date_span([last_day, spill_day]), html)
        self.assertNotIn('Girmemiş', html)
        self.assertNotIn('Tarihi gelmemiş', html)
        self.assertNotIn('>Oturum<', html)
        self.assertIn('SINAV TAKVİMİ', html)
        self.assertEqual(html.count('class="mark reg"'), 2)
        self.assertEqual(html.count('class="mark none"'), 2)
        from calendar import monthrange
        if today.day < monthrange(today.year, today.month)[1]:
            Exam.objects.create(
                name='Yarınki LGS', exam_type='LGS',
                exam_date=today + timedelta(days=1),
                kurum=self.kurum, sube=self.sube, egitim_yili=self.yil,
            )
            ctx, reason = load_giris_context(self.kurum.id, self.sube.id, self.exam.id, ogrenci.id)
            html = build_giris_belgesi_html(ctx)
            self.assertIn('YARINKİ LGS', html)
            chunk = html[html.index('YARINKİ LGS'):html.index('YARINKİ LGS') + 240]
            self.assertIn('class="mark none"', chunk)
            self.assertNotIn('class="mark reg"', chunk)

    def test_gecmis_ve_uyumsuz_deneme_kaydi_reddedilir(self):
        past = Exam.objects.create(
            name='Eski LGS', exam_type='LGS',
            exam_date=timezone.localdate() - timedelta(days=3),
            kurum=self.kurum, sube=self.sube, egitim_yili=self.yil,
        )
        self.exam = past
        payload = self._walk_payload()
        res = self._register(payload)
        self.assertEqual(res.status_code, 400, res.content)
        self.assertIn('geçmiş', res.json()['detail'])

        tyt = Exam.objects.create(
            name='TYT Deneme', exam_type='YKS_TYT',
            exam_date=timezone.localdate() + timedelta(days=8),
            kurum=self.kurum, sube=self.sube, egitim_yili=self.yil,
        )
        payload['gunu_birlik'] = {'exam_id': tyt.id}
        res = self._register(payload)
        self.assertEqual(res.status_code, 400, res.content)
        self.assertIn('sınıf seviyesine', res.json()['detail'])

    def test_deneme_ucreti_son_tutari_hatirlar_ve_bos_birakilirsa_yazmaz(self):
        from apps.finans.application.dashboard_overview_service import _bugun_alinan_toplam
        from apps.finans.domain.gelir_kaydi import GelirKaydi
        from apps.finans.domain.gelir_tahsilat import GelirTahsilat
        from apps.finans.infrastructure.bakiye_hareketi_repository import BakiyeHareketiRepository
        from apps.odeme_takip.domain.models import Sozlesme, Tahsilat

        hesap, yontem = self._nakit()
        self.exam = Exam.objects.create(
            name='LGS Deneme', exam_type='LGS', exam_date=timezone.localdate() + timedelta(days=12),
            kurum=self.kurum, sube=self.sube, egitim_yili=self.yil,
        )
        payload = self._walk_payload()
        payload['package'] = {}
        payload['gunu_birlik']['deneme_ucreti'] = '1.500'
        payload['gunu_birlik']['odeme_yontemi_id'] = yontem.id
        payload['gunu_birlik']['mali_hesap_id'] = hesap.id
        res = self._register(payload)
        self.assertEqual(res.status_code, 201, res.content)
        self.assertEqual(res.json()['deneme_ucreti'], 1500)
        self.assertNotIn('Ödeme şekli', res.json()['message'])
        ogrenci = Ogrenci.objects.get(tc_kimlik_no='10000000046')
        sozlesme = Sozlesme.objects.get(ogrenci=ogrenci)
        self.assertEqual(sozlesme.net_tutar, 1500)
        self.assertEqual(sozlesme.paket_turu, 'deneme')
        self.assertEqual(sozlesme.mali_hesap_id, hesap.id)
        self.assertEqual(sozlesme.odeme_yontemi_id, yontem.id)
        self.assertFalse(Tahsilat.objects.filter(sozlesme=sozlesme).exists())
        gelir = GelirKaydi.objects.get(aciklama__contains='deneme ücreti')
        self.assertEqual(gelir.durum, 'tahsil_edildi')
        self.assertEqual(int(gelir.tahsil_edilen), 1500)
        self.assertEqual(int(GelirTahsilat.objects.get(gelir_kaydi=gelir).tutar), 1500)
        self.assertEqual(int(BakiyeHareketiRepository.son_bakiye(hesap.id)), 1500)
        self.assertEqual(
            _bugun_alinan_toplam(self.kurum.id, timezone.localdate(), self.sube.id),
            1500,
        )

        listed = self.client.get('/api/ogrenci-kayit/gunu-birlik/', **self.headers)
        self.assertEqual(listed.status_code, 200, listed.content)
        row = next(item for item in listed.json()['ogrenciler'] if item['id'] == ogrenci.id)
        self.assertEqual(row['ogrenci_no'], 'GB-46')

        remembered = self.client.get(
            '/api/ogrenci-kayit/gunu-birlik/ucret/?tc=10000000046',
            **self.headers,
        )
        self.assertEqual(remembered.status_code, 200, remembered.content)
        self.assertEqual(remembered.json()['tutar'], 1500)

        empty = self._walk_payload(
            tc_kimlik_no='10000000050', ad='Bos', soyad='Ucret', telefon='05321110050',
        )
        empty['package'] = {}
        empty['enrollment']['ogrenci_no'] = 'GB-50'
        empty['gunu_birlik']['deneme_ucreti'] = ''
        res = self._register(empty)
        self.assertEqual(res.status_code, 201, res.content)
        bos = Ogrenci.objects.get(tc_kimlik_no='10000000050')
        self.assertFalse(Sozlesme.objects.filter(ogrenci=bos).exists())

        bad = self._walk_payload(
            tc_kimlik_no='10000000051', ad='Hatali', soyad='Ucret', telefon='05321110051',
        )
        bad['package'] = {}
        bad['gunu_birlik']['deneme_ucreti'] = '12,50'
        res = self._register(bad)
        self.assertEqual(res.status_code, 400, res.content)
        self.assertIn('tam sayı', res.json()['detail'])

    def test_oturum_degisir_iptal_finansi_kapatir_kurum_kaydi_listede_kalir(self):
        from apps.coaching.olcme_degerlendirme.models import ExamParticipant, ExamSessionModel
        from apps.odeme_takip.domain.models import Sozlesme, Tahsilat

        self.exam = Exam.objects.create(
            name='LGS Deneme', exam_type='LGS', exam_date=timezone.localdate() + timedelta(days=12),
            kurum=self.kurum, sube=self.sube, egitim_yili=self.yil,
        )
        hesap, yontem = self._nakit()
        payload = self._walk_payload(telefon='05321110060')
        payload['package'] = {}
        payload['gunu_birlik']['deneme_ucreti'] = '750'
        payload['gunu_birlik']['odeme_yontemi_id'] = yontem.id
        payload['gunu_birlik']['mali_hesap_id'] = hesap.id
        res = self._register(payload)
        self.assertEqual(res.status_code, 201, res.content)
        ogrenci = Ogrenci.objects.get(tc_kimlik_no='10000000046')
        sabah = ExamSessionModel.objects.create(
            exam=self.exam, name='Sabah', order=1,
            session_date=timezone.localdate() + timedelta(days=12),
        )
        aksam = ExamSessionModel.objects.create(
            exam=self.exam, name='Akşam', order=2,
            session_date=timezone.localdate() + timedelta(days=13),
        )
        moved = self.client.post(
            '/api/ogrenci-kayit/gunu-birlik/sinav/',
            data=json.dumps({
                'islem': 'oturum',
                'ogrenci_id': ogrenci.id,
                'exam_id': self.exam.id,
                'exam_session_id': aksam.id,
            }),
            content_type='application/json',
            **self.headers,
        )
        self.assertEqual(moved.status_code, 200, moved.content)
        participant = ExamParticipant.objects.get(exam=self.exam, student=ogrenci)
        self.assertEqual(participant.exam_session_id, aksam.id)
        self.assertNotEqual(participant.exam_session_id, sabah.id)
        self.assertTrue(Sozlesme.objects.filter(ogrenci=ogrenci, net_tutar=750).exclude(durum='iptal').exists())

        cancelled = self.client.post(
            '/api/ogrenci-kayit/gunu-birlik/sinav/',
            data=json.dumps({
                'islem': 'iptal',
                'ogrenci_id': ogrenci.id,
                'exam_id': self.exam.id,
            }),
            content_type='application/json',
            **self.headers,
        )
        self.assertEqual(cancelled.status_code, 200, cancelled.content)
        self.assertFalse(ExamParticipant.objects.filter(exam=self.exam, student=ogrenci).exists())
        from apps.finans.domain.gelir_kaydi import GelirKaydi
        from apps.finans.infrastructure.bakiye_hareketi_repository import BakiyeHareketiRepository

        sozlesme = Sozlesme.objects.get(ogrenci=ogrenci)
        self.assertEqual(sozlesme.durum, 'iptal')
        self.assertFalse(Tahsilat.objects.filter(sozlesme=sozlesme, durum='aktif').exists())
        self.assertEqual(GelirKaydi.objects.get(aciklama__contains='deneme ücreti').durum, 'iptal')
        self.assertEqual(int(BakiyeHareketiRepository.son_bakiye(hesap.id)), 0)

        from ..application.deneme_ucret import ensure_deneme_izi

        ogrenci.kayit_turu = 'asil'
        ogrenci.save(update_fields=['kayit_turu'])
        kayit = OgrenciKayit.objects.get(ogrenci=ogrenci)
        ensure_deneme_izi(ogrenci=ogrenci, kayit=kayit)
        listed = self.client.get('/api/ogrenci-kayit/gunu-birlik/', **self.headers)
        self.assertEqual(listed.status_code, 200, listed.content)
        self.assertTrue(any(item['id'] == ogrenci.id and item['kurum_kaydi'] for item in listed.json()['ogrenciler']))
        asil = Ogrenci.objects.create(
            kurum=self.kurum, sube=self.sube, ad='Duz', soyad='Ogrenci',
            tc_kimlik_no='10000000061', kayit_turu='asil', aktif_mi=True,
        )
        ExamParticipant.objects.create(
            exam=self.exam, student=asil, source=ExamParticipant.Source.MANUAL,
            sinif_seviyesi=self.seviye,
        )
        listed = self.client.get('/api/ogrenci-kayit/gunu-birlik/', **self.headers)
        self.assertFalse(any(item['id'] == asil.id for item in listed.json()['ogrenciler']))
        ctx = {
            'kurum_id': self.kurum.id,
            'sube_id': self.sube.id,
            'egitim_yili_id': self.yil.id,
        }
        params = {
            'q': '', 'all_years': False, 'durum': 'aktif',
            'sinif_seviyesi_ids': [], 'giris_turu': None, 'kayit_turu': None,
            'cinsiyet': None, 'paket_id': None, 'paket_turu': None, 'kalemler': [],
            'sinif_ids': [], 'school_ids': [], 'alan_ids': [], 'coach_ids': [],
            'kayit_tarihi_bas': None, 'kayit_tarihi_bit': None,
            'sort': 'created_at_desc', 'page': 1, 'page_size': 25,
        }
        normal, _ = build_kayit_queryset(ctx, params)
        self.assertIn(ogrenci.id, set(normal.values_list('ogrenci_id', flat=True)))
