"""Günü birlik deneme kaydı.

Öğrenci açılır, yıllık kayıt / sözleşme / paket açılmaz.
Katılım yalnız seçilen sınava yazılır.
"""
from __future__ import annotations

import logging
import re
from dataclasses import dataclass, field
from io import BytesIO

from django.db import transaction
from django.utils import timezone

logger = logging.getLogger(__name__)

from apps.egitim_tanimlari.models import SinifSeviyesi
from apps.kimlik.application.enforcement import assert_identity_for_new_record
from apps.kimlik.domain.phone import normalize_phone
from apps.kimlik.exceptions import KimlikConflictError
from apps.ogrenci.domain.models import Ogrenci, OgrenciKayit
from apps.ogrenci.services.kayit_turu import GUNU_BIRLIK_CODE

_HEADER_MAP = {
    'ad': 'ad',
    'soyad': 'soyad',
    'tc': 'tc',
    'tckimlik': 'tc',
    'tckimlikno': 'tc',
    'tcno': 'tc',
    'telefon': 'telefon',
    'tel': 'telefon',
    'gsm': 'telefon',
    'seviye': 'seviye',
    'sinifseviyesi': 'seviye',
    'sinif': 'seviye',
}


def _fold(value: str) -> str:
    table = str.maketrans('Iİıi', 'iiii')
    return (value or '').translate(table).casefold().strip()


def _compact(value: str) -> str:
    return re.sub(r'[^a-z0-9]', '', _fold(value))


def normalize_tc(value) -> str:
    return re.sub(r'\D', '', str(value or ''))


@dataclass
class RowOutcome:
    index: int
    ad: str = ''
    soyad: str = ''
    tc_kimlik_no: str = ''
    telefon: str = ''
    sinif_seviyesi_id: int | None = None
    sinif_seviyesi: str = ''
    status: str = ''
    message: str = ''
    ogrenci_id: int | None = None
    girilen_ad: str = ''
    girilen_soyad: str = ''
    _seviye: object = field(default=None, repr=False)

    def as_dict(self) -> dict:
        return {
            'index': self.index,
            'ad': self.ad,
            'soyad': self.soyad,
            'tc_kimlik_no': self.tc_kimlik_no,
            'telefon': self.telefon,
            'sinif_seviyesi_id': self.sinif_seviyesi_id,
            'sinif_seviyesi': self.sinif_seviyesi,
            'status': self.status,
            'message': self.message,
            'ogrenci_id': self.ogrenci_id,
            'girilen_ad': self.girilen_ad,
            'girilen_soyad': self.girilen_soyad,
        }


def load_seviyeler(kurum_id: int, sube_id: int):
    return list(
        SinifSeviyesi.objects.filter(
            kurum_id=kurum_id, sube_id=sube_id, aktif_mi=True,
        ).order_by('sira', 'ad')
    )


def _seviye_index(seviyeler):
    by_id = {s.id: s for s in seviyeler}
    by_key = {}
    for seviye in seviyeler:
        for raw in (seviye.kod, seviye.ad):
            folded = _fold(raw)
            compact = _compact(raw)
            by_key.setdefault(folded, seviye)
            if compact:
                by_key.setdefault(compact, seviye)
    return by_id, by_key


def _resolve_seviye(raw, by_id, by_key):
    if raw in (None, ''):
        return None
    if isinstance(raw, int) or (isinstance(raw, str) and str(raw).strip().isdigit()):
        found = by_id.get(int(raw))
        if found:
            return found
    text = str(raw).strip()
    return by_key.get(_fold(text)) or by_key.get(_compact(text))


def _cell(value) -> str:
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_excel(uploaded) -> list[dict]:
    from openpyxl import load_workbook

    wb = load_workbook(uploaded, read_only=True, data_only=True)
    ws = wb.active
    rows = []
    for values in ws.iter_rows(values_only=True):
        if not values or all(v in (None, '') for v in values):
            continue
        rows.append([_cell(v) for v in values])
    wb.close()
    if not rows:
        return []

    header_keys = [_compact(c) for c in rows[0]]
    mapped = [_HEADER_MAP.get(k) for k in header_keys]
    has_header = 'ad' in mapped and 'soyad' in mapped and 'tc' in mapped
    data_rows = rows[1:] if has_header else rows
    if not has_header:
        mapped = ['ad', 'soyad', 'tc', 'telefon', 'seviye']

    parsed = []
    for raw in data_rows:
        item = {'ad': '', 'soyad': '', 'tc': '', 'telefon': '', 'seviye': ''}
        for idx, key in enumerate(mapped):
            if not key or idx >= len(raw):
                continue
            item[key] = raw[idx]
        if not any(item.values()):
            continue
        parsed.append({
            'ad': item['ad'],
            'soyad': item['soyad'],
            'tc_kimlik_no': item['tc'],
            'telefon': item['telefon'],
            'sinif_seviyesi': item['seviye'],
        })
    return parsed


def excel_template() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = 'Tek deneme'
    ws.append(['Ad', 'Soyad', 'TC', 'Telefon', 'Seviye'])
    ws.append(['Ayşe', 'Yılmaz', '12345678901', '05321112233', '12'])
    ws.append([])
    ws.append(['Seviye sütununa sınıf seviyesi kodu veya adı yazılır (ör. 12 veya 12. Sınıf).'])
    for col, width in zip('ABCDE', (18, 18, 16, 16, 18)):
        ws.column_dimensions[col].width = width
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _blank_outcome(index: int, raw: dict) -> RowOutcome:
    return RowOutcome(
        index=index,
        ad=(raw.get('ad') or '').strip(),
        soyad=(raw.get('soyad') or '').strip(),
        tc_kimlik_no=normalize_tc(raw.get('tc_kimlik_no') or raw.get('tc')),
        telefon=normalize_phone(raw.get('telefon') or ''),
        girilen_ad=(raw.get('ad') or '').strip(),
        girilen_soyad=(raw.get('soyad') or '').strip(),
    )


def classify_rows(exam, raw_rows, seviyeler, *, exam_session_id=None) -> list[RowOutcome]:
    from apps.coaching.olcme_degerlendirme.models import ExamParticipant

    from .sinav_uygunluk import ExamLevelGate, upcoming_sessions

    by_id, by_key = _seviye_index(seviyeler)
    all_sessions = list(exam.exam_sessions.order_by('order', 'id'))
    sessions = upcoming_sessions(exam, sessions=all_sessions)
    if all_sessions and not sessions:
        session_error = 'Bu denemenin seçilebilir oturumu kalmadı.'
    else:
        session_error = _session_error(sessions, exam_session_id)
    level_gate = ExamLevelGate(exam)

    outcomes: list[RowOutcome] = []
    seen: set[str] = set()
    tcs = []
    for index, raw in enumerate(raw_rows or []):
        if not isinstance(raw, dict):
            raw = {}
        row = _blank_outcome(index, raw)
        seviye_raw = raw.get('sinif_seviyesi_id')
        if seviye_raw in (None, ''):
            seviye_raw = raw.get('sinif_seviyesi') or raw.get('seviye')
        seviye = _resolve_seviye(seviye_raw, by_id, by_key)
        if seviye:
            row.sinif_seviyesi_id = seviye.id
            row.sinif_seviyesi = seviye.ad
            row._seviye = seviye
        errors = []
        if not row.ad:
            errors.append('Ad zorunlu.')
        if not row.soyad:
            errors.append('Soyad zorunlu.')
        if len(row.tc_kimlik_no) != 11:
            errors.append('TC 11 hane olmalı.')
        elif row.tc_kimlik_no in seen:
            errors.append('Bu TC listede tekrar ediyor.')
        else:
            seen.add(row.tc_kimlik_no)
            tcs.append(row.tc_kimlik_no)
        if not seviye:
            errors.append('Sınıf seviyesi bulunamadı.')
        elif not level_gate.matches(seviye):
            errors.append('Bu deneme seçilen sınıf seviyesine uygun değil.')
        if session_error:
            errors.append(session_error)
        if errors:
            row.status = 'hata'
            row.message = ' '.join(errors)
        outcomes.append(row)

    if not tcs:
        return outcomes

    existing = {
        o.tc_kimlik_no: o
        for o in Ogrenci.objects.filter(kurum_id=exam.kurum_id, tc_kimlik_no__in=tcs)
    }
    session = _target_session(sessions, exam_session_id)
    participant_qs = ExamParticipant.objects.filter(exam=exam, student__tc_kimlik_no__in=tcs)
    if session is not None:
        participant_qs = participant_qs.filter(exam_session=session)
    elif not sessions:
        participant_qs = participant_qs.filter(exam_session__isnull=True)
    on_exam = set(participant_qs.values_list('student__tc_kimlik_no', flat=True))

    for row in outcomes:
        if row.status == 'hata' or not row.tc_kimlik_no:
            continue
        ogrenci = existing.get(row.tc_kimlik_no)
        if ogrenci is None:
            row.status = 'yeni'
            row.message = 'Yeni tek deneme kaydı açılacak.'
            continue
        row.ogrenci_id = ogrenci.id
        row.ad = ogrenci.ad
        row.soyad = ogrenci.soyad
        if ogrenci.sube_id != exam.sube_id:
            row.status = 'hata'
            row.message = 'Bu TC başka şubede kayıtlı.'
            continue
        if not ogrenci.aktif_mi:
            row.status = 'hata'
            row.message = 'Öğrenci pasif.'
            continue
        if row.tc_kimlik_no in on_exam:
            row.status = 'zaten_sinavda'
            row.message = 'Bu öğrenci zaten bu sınavda.'
            continue
        if (ogrenci.kayit_turu or '') == GUNU_BIRLIK_CODE:
            row.status = 'mevcut_misafir'
            row.message = 'Mevcut tek deneme kaydı bu sınava eklenecek.'
        else:
            row.status = 'kayitli_ogrenci'
            row.message = 'Kayıtlı öğrenci, sınava eklenecek.'
    return outcomes


def _session_error(sessions, exam_session_id) -> str:
    if len(sessions) <= 1:
        if exam_session_id and not any(s.pk == exam_session_id for s in sessions):
            return 'Oturum bulunamadı.'
        return ''
    if not exam_session_id:
        return 'Oturum seçin.'
    if not any(s.pk == exam_session_id for s in sessions):
        return 'Oturum bulunamadı.'
    return ''


def _target_session(sessions, exam_session_id):
    if exam_session_id:
        return next((s for s in sessions if s.pk == exam_session_id), None)
    if len(sessions) == 1:
        return sessions[0]
    return None


def _attach(exam, student, seviye_id, exam_session_id):
    from apps.coaching.olcme_degerlendirme.models import ExamParticipant
    from apps.coaching.olcme_degerlendirme.services.exam_roster import add_manual_participant

    from .sinav_uygunluk import upcoming_sessions

    all_sessions = list(exam.exam_sessions.order_by('order', 'id'))
    sessions = upcoming_sessions(exam, sessions=all_sessions)
    if all_sessions and not sessions:
        return None, 'Bu denemenin seçilebilir oturumu kalmadı.'
    err = _session_error(sessions, exam_session_id)
    if err:
        return None, err
    target = _target_session(sessions, exam_session_id)
    qs = ExamParticipant.objects.filter(exam=exam, student=student)
    qs = qs.filter(exam_session=target) if target is not None else qs.filter(exam_session__isnull=True)
    current = qs.first()
    if current:
        _stamp_seviye(current, student, seviye_id)
        return current, 'zaten'

    participant, add_err = add_manual_participant(exam, student.id, exam_session=target)
    if add_err and participant is None:
        return None, add_err
    if participant is not None:
        _stamp_seviye(participant, student, seviye_id)
    if add_err and 'zaten' in add_err:
        return participant, 'zaten'
    return participant, None


def _stamp_seviye(participant, student, seviye_id):
    if (student.kayit_turu or '') != GUNU_BIRLIK_CODE or not seviye_id:
        return
    if participant.sinif_seviyesi_id == seviye_id:
        return
    participant.sinif_seviyesi_id = seviye_id
    participant.save(update_fields=['sinif_seviyesi', 'updated_at'])


def commit_rows(exam, raw_rows, seviyeler, *, exam_session_id=None) -> list[RowOutcome]:
    preview = classify_rows(exam, raw_rows, seviyeler, exam_session_id=exam_session_id)
    saved: list[RowOutcome] = []
    for row in preview:
        if row.status == 'hata':
            saved.append(row)
            continue
        try:
            with transaction.atomic():
                _apply_row(exam, row, exam_session_id=exam_session_id)
        except KimlikConflictError as exc:
            row.status = 'hata'
            row.message = exc.message
        except Exception as exc:
            row.status = 'hata'
            row.message = str(exc) or 'Kayıt yazılamadı.'
        saved.append(row)
    return saved


def _apply_row(exam, row: RowOutcome, *, exam_session_id=None):
    ogrenci = None
    if row.ogrenci_id:
        ogrenci = Ogrenci.objects.select_for_update().filter(pk=row.ogrenci_id).first()
    if ogrenci is None and row.tc_kimlik_no:
        ogrenci = (
            Ogrenci.objects.select_for_update()
            .filter(kurum_id=exam.kurum_id, tc_kimlik_no=row.tc_kimlik_no)
            .first()
        )

    if ogrenci is None:
        assert_identity_for_new_record(
            exam.kurum_id,
            tc_kimlik_no=row.tc_kimlik_no,
            telefon='',
            context='ogrenci',
        )
        ogrenci = Ogrenci.objects.create(
            kurum_id=exam.kurum_id,
            sube_id=exam.sube_id,
            tc_kimlik_no=row.tc_kimlik_no,
            ad=row.girilen_ad or row.ad,
            soyad=row.girilen_soyad or row.soyad,
            telefon=row.telefon,
            kayit_turu=GUNU_BIRLIK_CODE,
            aktif_mi=True,
        )
        row.status = 'yeni'
        row.message = 'Tek deneme kaydı açıldı ve sınava eklendi.'
    else:
        if ogrenci.sube_id != exam.sube_id:
            row.status = 'hata'
            row.message = 'Bu TC başka şubede kayıtlı.'
            transaction.set_rollback(True)
            return
        if not ogrenci.aktif_mi:
            row.status = 'hata'
            row.message = 'Öğrenci pasif.'
            transaction.set_rollback(True)
            return
        if (ogrenci.kayit_turu or '') == GUNU_BIRLIK_CODE:
            if row.telefon and ogrenci.telefon != row.telefon:
                ogrenci.telefon = row.telefon
                ogrenci.save(update_fields=['telefon', 'updated_at'])
            row.status = 'mevcut_misafir'
            row.message = 'Mevcut tek deneme kaydı bu sınava eklendi.'
        else:
            row.status = 'kayitli_ogrenci'
            row.message = 'Kayıtlı öğrenci, sınava eklendi.'

    row.ogrenci_id = ogrenci.id
    row.ad = ogrenci.ad
    row.soyad = ogrenci.soyad
    participant, err = _attach(exam, ogrenci, row.sinif_seviyesi_id, exam_session_id)
    if err == 'zaten':
        row.status = 'zaten_sinavda'
        row.message = 'Bu öğrenci zaten bu sınavda.'
        return
    if err:
        row.status = 'hata'
        row.message = err
        transaction.set_rollback(True)
        return
    if participant is None:
        row.status = 'hata'
        row.message = 'Sınav listesine eklenemedi.'
        transaction.set_rollback(True)


def summarize(rows: list[RowOutcome]) -> dict:
    counts = {
        'yeni': 0,
        'mevcut_misafir': 0,
        'kayitli_ogrenci': 0,
        'zaten_sinavda': 0,
        'hata': 0,
    }
    for row in rows:
        counts[row.status] = counts.get(row.status, 0) + 1
    return counts


def assert_no_yearly_enrollment(ogrenci_id: int) -> bool:
    return not OgrenciKayit.objects.filter(ogrenci_id=ogrenci_id).exists()


def _deneme_izi_var():
    from django.db.models import Exists, OuterRef

    from apps.odeme_takip.domain.models import Sozlesme

    from .deneme_ucret import MARKER

    return Exists(
        Sozlesme.objects.filter(ogrenci_id=OuterRef('pk'), notlar__contains=MARKER)
    )


def list_misafirler(kurum_id: int, sube_id: int, *, q: str = '') -> list[dict]:
    from django.db.models import Q

    from apps.coaching.olcme_degerlendirme.models import ExamParticipant

    from .sinav_uygunluk import upcoming_sessions

    today = timezone.localdate()
    qs = Ogrenci.objects.filter(
        kurum_id=kurum_id,
        sube_id=sube_id,
        aktif_mi=True,
    ).filter(
        Q(kayit_turu=GUNU_BIRLIK_CODE) | _deneme_izi_var()
    ).order_by('-created_at', 'soyad', 'ad')
    text = (q or '').strip()
    if text:
        from django.db.models import Q
        qs = qs.filter(
            Q(ad__icontains=text)
            | Q(soyad__icontains=text)
            | Q(tc_kimlik_no__icontains=text)
            | Q(telefon__icontains=normalize_tc(text) or text)
        )
    ogrenciler = list(qs[:500])
    ids = [o.id for o in ogrenciler]
    numaralar = {}
    seviye_ids = {}
    seviye_adlari = {}
    for kayit in (
        OgrenciKayit.objects.filter(ogrenci_id__in=ids)
        .select_related('sinif_seviyesi')
        .order_by('id')
    ):
        numaralar[kayit.ogrenci_id] = kayit.okul_no or ''
        if kayit.sinif_seviyesi_id:
            seviye_ids[kayit.ogrenci_id] = kayit.sinif_seviyesi_id
            seviye_adlari[kayit.ogrenci_id] = kayit.sinif_seviyesi.ad if kayit.sinif_seviyesi_id else ''
    sinavlar: dict[int, list] = {i: [] for i in ids}
    seen_pairs: set[tuple[int, int]] = set()
    parts = (
        ExamParticipant.objects.filter(student_id__in=ids)
        .select_related('exam', 'exam_session', 'sinif_seviyesi')
        .prefetch_related('exam__exam_sessions')
        .order_by('-exam__exam_date', '-exam_id', 'id')
    )
    session_cache: dict[int, list] = {}
    for part in parts:
        key = (part.student_id, part.exam_id)
        if key in seen_pairs or not part.exam_id:
            continue
        seen_pairs.add(key)
        exam = part.exam
        if exam.id not in session_cache:
            all_sessions = list(exam.exam_sessions.all())
            upcoming = upcoming_sessions(exam, today=today, sessions=all_sessions)
            open_exam = bool(upcoming) if all_sessions else bool(exam.exam_date and exam.exam_date >= today)
            session_cache[exam.id] = (upcoming, open_exam)
        upcoming, open_exam = session_cache[exam.id]
        dated = [session.session_date for session in upcoming if session.session_date]
        bookable = min(dated) if dated else exam.exam_date
        current = part.exam_session
        seviye = part.sinif_seviyesi.ad if part.sinif_seviyesi_id else ''
        sinavlar[part.student_id].append({
            'id': exam.id,
            'participant_id': part.id,
            'name': exam.name,
            'exam_date': bookable.isoformat() if bookable else (
                exam.exam_date.isoformat() if exam.exam_date else None
            ),
            'sinif_seviyesi': seviye,
            'acik': open_exam,
            'session_id': current.id if current is not None else None,
            'session_name': current.name if current is not None else '',
            'sessions': [
                {
                    'id': session.id,
                    'name': session.name,
                    'session_date': session.session_date.isoformat() if session.session_date else None,
                }
                for session in upcoming
            ],
        })
    return [
        {
            'id': ogrenci.id,
            'ad': ogrenci.ad,
            'soyad': ogrenci.soyad,
            'tc_kimlik_no': ogrenci.tc_kimlik_no or '',
            'telefon': ogrenci.telefon or '',
            'ogrenci_no': numaralar.get(ogrenci.id) or '',
            'sinif_seviyesi_id': seviye_ids.get(ogrenci.id),
            'sinif_seviyesi': seviye_adlari.get(ogrenci.id) or '',
            'kayit_turu': ogrenci.kayit_turu or '',
            'kurum_kaydi': (ogrenci.kayit_turu or '') != GUNU_BIRLIK_CODE,
            'deneme_sayisi': len(sinavlar.get(ogrenci.id, [])),
            'created_at': ogrenci.created_at.isoformat() if ogrenci.created_at else None,
            'sinavlar': sinavlar.get(ogrenci.id, []),
        }
        for ogrenci in ogrenciler
    ]


def _deneme_exam_ogrenci(kurum_id, sube_id, exam_id, ogrenci_id):
    from apps.coaching.olcme_degerlendirme.models import Exam, ExamParticipant

    from .sinav_uygunluk import exam_is_open

    exam = Exam.objects.filter(id=exam_id, kurum_id=kurum_id, sube_id=sube_id, is_active=True).first()
    if exam is None:
        raise ValueError('Sınav bulunamadı.')
    if not exam_is_open(exam):
        raise ValueError('Tarihi geçmiş denemede bu işlem yapılamaz.')
    ogrenci = Ogrenci.objects.filter(
        id=ogrenci_id, kurum_id=kurum_id, sube_id=sube_id, aktif_mi=True,
    ).first()
    if ogrenci is None:
        raise ValueError('Öğrenci bulunamadı.')
    participant = (
        ExamParticipant.objects.filter(exam=exam, student=ogrenci)
        .select_related('exam_session')
        .order_by('id')
        .first()
    )
    if participant is None:
        raise ValueError('Bu öğrencinin bu sınavda kaydı yok.')
    return exam, ogrenci, participant


def register_deneme_kaydi(
    *,
    kurum_id,
    sube_id,
    ogrenci_id,
    exam_id,
    exam_session_id=None,
    deneme_ucreti=None,
    user=None,
    odeme_yontemi_id=None,
    mali_hesap_id=None,
) -> dict:
    """Listeden seçilen denemeye, kayıt sihirbazındaki gibi ekler."""
    from apps.coaching.olcme_degerlendirme.models import Exam

    from .deneme_ucret import record_deneme_ucreti
    from .sinav_uygunluk import exam_registration_error

    try:
        ogrenci_id = int(ogrenci_id)
        exam_id = int(exam_id)
    except (TypeError, ValueError):
        raise ValueError('Öğrenci ve sınav seçin.')
    if exam_session_id in ('', None):
        exam_session_id = None
    else:
        try:
            exam_session_id = int(exam_session_id)
        except (TypeError, ValueError):
            raise ValueError('Oturum bulunamadı.')

    exam = Exam.objects.filter(
        id=exam_id, kurum_id=kurum_id, sube_id=sube_id, is_active=True,
    ).first()
    if exam is None:
        raise ValueError('Sınav bulunamadı.')
    ogrenci = Ogrenci.objects.filter(
        id=ogrenci_id, kurum_id=kurum_id, sube_id=sube_id, aktif_mi=True,
    ).first()
    if ogrenci is None:
        raise ValueError('Öğrenci bulunamadı.')
    kayit = (
        OgrenciKayit.objects.filter(ogrenci=ogrenci)
        .select_related('sinif_seviyesi')
        .order_by('-id')
        .first()
    )
    seviye = kayit.sinif_seviyesi if kayit is not None and kayit.sinif_seviyesi_id else None
    if seviye is None:
        raise ValueError('Sınıf seviyesi bulunamadı.')
    registration_error = exam_registration_error(exam, seviye)
    if registration_error:
        raise ValueError(registration_error)

    with transaction.atomic():
        participant, attach_err = _attach(exam, ogrenci, seviye.id, exam_session_id)
        if attach_err == 'zaten':
            raise ValueError('Bu öğrenci zaten bu sınavda.')
        if attach_err:
            raise ValueError(attach_err)
        seat = assign_open_seat(participant) if participant is not None else {'status': '', 'detail': ''}
    ucret_uyari = ''
    if deneme_ucreti is not None:
        try:
            record_deneme_ucreti(
                ogrenci=ogrenci,
                kayit=kayit,
                exam=exam,
                tutar=deneme_ucreti,
                user=user,
                odeme_yontemi_id=odeme_yontemi_id,
                mali_hesap_id=mali_hesap_id,
            )
        except ValueError as exc:
            ucret_uyari = str(exc)
        except Exception:
            logger.exception('Deneme ücreti yazılamadı')
            ucret_uyari = 'Deneme ücreti kayda geçmedi.'
    detail = seat.get('detail') or ''
    if ucret_uyari:
        detail = f'{detail} {ucret_uyari}'.strip()
    return {
        'ok': True,
        'exam_id': exam.id,
        'sinav_adi': exam.name,
        'participant_id': participant.id if participant is not None else None,
        'seat_status': seat.get('status') or '',
        'seat_detail': detail,
        'deneme_ucreti': None if ucret_uyari else deneme_ucreti,
    }


def change_deneme_oturumu(*, kurum_id, sube_id, ogrenci_id, exam_id, exam_session_id) -> dict:
    from .sinav_uygunluk import upcoming_sessions

    exam, _ogrenci, participant = _deneme_exam_ogrenci(kurum_id, sube_id, exam_id, ogrenci_id)
    try:
        session_id = int(exam_session_id)
    except (TypeError, ValueError):
        raise ValueError('Oturum bulunamadı.')
    target = next((item for item in upcoming_sessions(exam) if item.id == session_id), None)
    if target is None:
        raise ValueError('Bu oturum artık seçilemez.')
    if participant.exam_session_id == target.id:
        return {'ok': True, 'session_id': target.id, 'session_name': target.name, 'seat_status': 'assigned'}
    from apps.coaching.olcme_degerlendirme.models import ExamParticipant

    clash = ExamParticipant.objects.filter(
        exam=exam, student_id=ogrenci_id, exam_session=target,
    ).exclude(pk=participant.pk).exists()
    if clash:
        raise ValueError('Öğrenci bu oturumda zaten kayıtlı.')
    participant.exam_session = target
    participant.room = None
    participant.seat_no = None
    participant.desk_no = ''
    participant.save(update_fields=['exam_session', 'room', 'seat_no', 'desk_no', 'updated_at'])
    seat = assign_open_seat(participant)
    return {
        'ok': True,
        'session_id': target.id,
        'session_name': target.name,
        'seat_status': seat.get('status') or '',
        'seat_detail': seat.get('detail') or '',
    }


def cancel_deneme_kaydi(*, kurum_id, sube_id, ogrenci_id, exam_id) -> dict:
    from django.db.models.deletion import ProtectedError

    from apps.coaching.olcme_degerlendirme.models import ExamParticipant

    from .deneme_ucret import void_deneme_ucreti

    exam, ogrenci, _participant = _deneme_exam_ogrenci(kurum_id, sube_id, exam_id, ogrenci_id)
    try:
        with transaction.atomic():
            ExamParticipant.objects.filter(exam=exam, student=ogrenci).delete()
            closed = void_deneme_ucreti(ogrenci=ogrenci, exam=exam)
    except ProtectedError:
        raise ValueError('Bu denemeye sonuç girildiği için kayıt iptal edilemedi.')
    return {'ok': True, 'finans_iptal': closed}


def assign_open_seat(participant) -> dict:
    """Boş salon sırası varsa yalnız bu katılımcıyı yerleştirir.

    Kayıt salon dolu diye düşmez. Sonuç kayıt cevabında ve belgede görünür.
    """
    if participant is None:
        return {'status': 'none', 'detail': ''}
    if participant.room_id and participant.seat_no:
        return {'status': 'assigned', 'detail': ''}
    try:
        from apps.coaching.olcme_degerlendirme.services.exam_roster import (
            assign_participant_to_room,
            rooms_for_session,
        )

        rooms = rooms_for_session(participant.exam, participant.exam_session)
        if not rooms:
            return {
                'status': 'no_room',
                'detail': 'Bu oturumda salon tanımlı değil. Salon ve sıra sınav günü duyurulacak.',
            }
        last_error = ''
        for room in rooms:
            last_error = assign_participant_to_room(participant, room) or ''
            if not last_error:
                participant.refresh_from_db()
                return {'status': 'assigned', 'detail': ''}
        return {
            'status': 'full',
            'detail': last_error or 'Salonda boş sıra kalmadı. Sıra sınav günü duyurulacak.',
        }
    except Exception:
        logger.exception('Günü birlik katılımcı sırası atanamadı: %s', getattr(participant, 'pk', None))
        return {
            'status': 'error',
            'detail': 'Salon sırası şu an atanamadı. Yer bilgisi sınav günü duyurulacak.',
        }


def exam_options(kurum_id: int, sube_id: int, egitim_yili_id: int | None, sinif_seviyesi_id: int | None = None):
    from apps.coaching.olcme_degerlendirme.models import Exam

    from .sinav_uygunluk import exam_is_open, exam_matches_seviye, target_seviye_ids, upcoming_sessions

    seviye = None
    if sinif_seviyesi_id:
        seviye = SinifSeviyesi.objects.filter(pk=sinif_seviyesi_id, aktif_mi=True).first()
        if seviye is None:
            return []

    qs = Exam.objects.filter(
        kurum_id=kurum_id, sube_id=sube_id, is_active=True,
    )
    if egitim_yili_id:
        qs = qs.filter(egitim_yili_id=egitim_yili_id)
    qs = qs.prefetch_related(
        'exam_sessions',
        'audiences__deneme_paketi__sinif_seviyeleri',
        'deneme_paketi__sinif_seviyeleri',
        'siniflar',
    ).order_by('exam_date', 'name', 'id')
    exams = []
    for exam in qs:
        if not exam_is_open(exam):
            continue
        if seviye is not None and not exam_matches_seviye(exam, seviye):
            continue
        sessions = upcoming_sessions(exam, sessions=list(exam.exam_sessions.all()))
        dated = [session.session_date for session in sessions if session.session_date]
        bookable = min(dated) if dated else exam.exam_date
        exams.append({
            'id': exam.id,
            'name': exam.name,
            'exam_date': bookable.isoformat() if bookable else None,
            'exam_type': exam.get_exam_type_display(),
            'status': exam.status,
            'sinif_seviyesi_ids': sorted(target_seviye_ids(exam)),
            'sessions': [
                {
                    'id': session.id,
                    'name': session.name,
                    'session_date': session.session_date.isoformat() if session.session_date else None,
                    'schedule_preference_display': session.get_schedule_preference_display(),
                }
                for session in sessions
            ],
        })
    return exams
