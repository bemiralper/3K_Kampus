"""Okulizyon'un 9–11. sınıf Maarif sayfalarını mevcut ders ağacına yazar.

Yalnızca sayfa adında "Maarif" geçen yapraklar okunur. 12-TYT-AYT ve AYT-DİL
yaprakları 2018 kataloğudur; onlara dokunulmaz.

Aynı dersin 9, 10 ve 11. sınıf Maarif konuları yenilenir. 2018 programı,
5–8 ve 12. sınıf Maarif satırları durur. Cevap anahtarı bağları silinmez;
konu kalkınca bağ boşalır (SET NULL).
"""
from __future__ import annotations

from collections import defaultdict
from pathlib import Path

from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.db.models import Q

from apps.coaching.olcme_degerlendirme.models.answer_key import AnswerKeyItem
from apps.coaching.olcme_degerlendirme.models.curriculum import (
    Outcome,
    Subject,
    SubOutcome,
    Topic,
)
from apps.coaching.olcme_degerlendirme.services.okulizyon_import import DERS_SUBJECTS

LISE_GRADES = frozenset({'9', '10', '11'})


def _parts(code: str) -> list[str]:
    return [part for part in str(code).split('.') if part]


def _parent(code: str) -> str:
    parts = _parts(code)
    return '.'.join(parts[:-1]) if len(parts) > 1 else code


def parse_maarif_files(paths: list[Path]) -> tuple[list[dict], dict]:
    from openpyxl import load_workbook

    rows: list[dict] = []
    seen: set[tuple[str, str]] = set()
    skipped = {'sheets': [], 'duplicate_rows': 0, 'blank_rows': 0, 'unknown_ders': {}}
    for path in paths:
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in wb.sheetnames:
                if 'maarif' not in sheet.casefold():
                    skipped['sheets'].append(f'{path.name} / {sheet}')
                    continue
                ws = wb[sheet]
                iterator = ws.iter_rows(values_only=True)
                header = next(iterator, None)
                if not header or 'kaz' not in str(header[3] or '').casefold():
                    raise CommandError(
                        f'{path.name} / {sheet}: beklenen sütun "Okulizyon Kaz. Kodu" değil.'
                    )
                for raw in iterator:
                    if not raw or all(cell is None or str(cell).strip() == '' for cell in raw[:5]):
                        skipped['blank_rows'] += 1
                        continue
                    ders = str(raw[1] or '').strip()
                    code = str(raw[3] or '').strip().rstrip('.')
                    text = str(raw[4] or '').strip()
                    grade = str(raw[2] or '').strip()
                    if not ders or not code:
                        skipped['blank_rows'] += 1
                        continue
                    if ders not in DERS_SUBJECTS:
                        skipped['unknown_ders'][ders] = skipped['unknown_ders'].get(ders, 0) + 1
                        continue
                    if grade not in LISE_GRADES:
                        raise CommandError(
                            f'{path.name} / {sheet}: Maarif sayfasında beklenmeyen sınıf {grade!r} ({code}).'
                        )
                    key = (DERS_SUBJECTS[ders]['code'], code)
                    if key in seen:
                        skipped['duplicate_rows'] += 1
                        continue
                    seen.add(key)
                    rows.append({
                        'ders': ders,
                        'subject_code': key[0],
                        'code': code,
                        'grade': grade,
                        'text': text or code,
                    })
        finally:
            wb.close()
    return rows, skipped


def rows_to_topics(rows: list[dict]) -> list[dict]:
    """2 parça ünite, 4 parça kazanım, 5 parça alt kazanım.

    3 parçalı satır ara başlıktır. Altında kazanım yoksa kendisi kazanım olur.
    Bu, mevcut Maarif kataloğundaki düzeyle aynıdır.
    """
    topics: list[dict] = []
    current_topic: dict | None = None
    current_outcome: dict | None = None
    pending: dict | None = None

    def ensure_topic(code: str, grade: str, text: str) -> dict:
        nonlocal current_topic, current_outcome, pending
        current_topic = {
            'code': code[:30],
            'name': f'Maarif · {grade}. sınıf · {text}'[:200],
            'outcomes': [],
        }
        topics.append(current_topic)
        current_outcome = None
        pending = None
        return current_topic

    def flush_pending() -> None:
        nonlocal pending, current_outcome
        if pending is None or current_topic is None:
            pending = None
            return
        current_outcome = {'code': pending['code'][:50], 'text': pending['text'], 'sub_outcomes': []}
        current_topic['outcomes'].append(current_outcome)
        pending = None

    for row in rows:
        code = row['code']
        parts = _parts(code)
        text = row['text']
        if len(parts) <= 1:
            continue
        if len(parts) == 2:
            flush_pending()
            ensure_topic(code, row['grade'], text)
            continue
        if current_topic is None:
            ensure_topic(_parent(code), row['grade'], _parent(code))
        if len(parts) == 3:
            flush_pending()
            pending = {'code': code, 'text': text}
            current_outcome = None
            continue
        if len(parts) == 4:
            pending = None
            current_outcome = {'code': code[:50], 'text': text, 'sub_outcomes': []}
            current_topic['outcomes'].append(current_outcome)
            continue
        if current_outcome is None:
            current_outcome = {
                'code': _parent(code)[:50],
                'text': _parent(code),
                'sub_outcomes': [],
            }
            current_topic['outcomes'].append(current_outcome)
        current_outcome['sub_outcomes'].append({'code': code[:50], 'text': text})
    flush_pending()
    return topics


def _protected_topic_ids(topic_ids: list[int]) -> set[int]:
    from apps.mobile_identity.models import IslenenKonu

    if not topic_ids:
        return set()
    return set(
        IslenenKonu.objects.filter(topic_id__in=topic_ids).values_list('topic_id', flat=True)
    )


def _write_outcomes(topic: Topic, outcomes: list[dict]) -> None:
    """Korumalı kazanım satırını yerinde günceller, kalanını yeniler."""
    from apps.mobile_identity.models import IslenenKazanim

    protected = {
        row.code: row
        for row in Outcome.objects.filter(
            topic=topic,
            id__in=IslenenKazanim.objects.filter(outcome__topic=topic).values('outcome_id'),
        )
    }
    topic.outcomes.exclude(id__in=[row.id for row in protected.values()]).delete()
    seen = set(protected)
    for index, outcome_data in enumerate(outcomes):
        code = outcome_data['code']
        existing = protected.get(code)
        if existing is None and code in seen:
            continue
        if existing is None:
            existing = Outcome.objects.create(
                topic=topic, code=code, text=outcome_data['text'], order=index,
            )
        else:
            existing.text = outcome_data['text']
            existing.order = index
            existing.save(update_fields=['text', 'order'])
        seen.add(code)
        existing.sub_outcomes.all().delete()
        if outcome_data['sub_outcomes']:
            SubOutcome.objects.bulk_create([
                SubOutcome(
                    outcome=existing, code=sub['code'], text=sub['text'], order=sub_index,
                )
                for sub_index, sub in enumerate(outcome_data['sub_outcomes'])
            ])


def _grade_filter(grades: set[str]) -> Q:
    grade_re = '|'.join(sorted(grades, key=int))
    return (
        Q(code__regex=rf'^({grade_re})\.')
        | Q(name__regex=rf'(^|[·•])\s*({grade_re})\. sınıf')
    )


class Command(BaseCommand):
    help = 'Okulizyon 9–11. sınıf Maarif Excel sayfalarını ders ağacına yükler.'

    def add_arguments(self, parser):
        parser.add_argument('paths', nargs='+', help='xlsx dosyası ya da klasör')
        parser.add_argument('--dry-run', action='store_true', help='Yazmadan say.')

    def handle(self, *args, **options):
        files: list[Path] = []
        for raw in options['paths']:
            path = Path(raw).expanduser()
            if path.is_dir():
                files.extend(sorted(p for p in path.glob('*.xlsx') if not p.name.startswith('~$')))
            elif path.is_file():
                files.append(path)
            else:
                raise CommandError(f'Bulunamadı: {path}')
        if not files:
            raise CommandError('Excel dosyası yok.')

        rows, skipped = parse_maarif_files(files)
        by_subject: dict[str, list[dict]] = defaultdict(list)
        for row in rows:
            by_subject[row['subject_code']].append(row)

        self.stdout.write(
            f'{len(files)} dosya, {len(rows)} Maarif satırı, '
            f'{len(by_subject)} ders. Yinelenen satır: {skipped["duplicate_rows"]}.'
        )
        if skipped['unknown_ders']:
            raise CommandError(f'Tanınmayan ders adı: {skipped["unknown_ders"]}')
        if skipped['sheets']:
            self.stdout.write('Atlanan yapraklar (2018 / dil, yüklenmedi):')
            for label in skipped['sheets']:
                self.stdout.write(f'  - {label}')

        plan = []
        for code, subject_rows in by_subject.items():
            spec = next(item for item in DERS_SUBJECTS.values() if item['code'] == code)
            subject = Subject.objects.filter(code=code).first()
            grades = {row['grade'] for row in subject_rows}
            topics = rows_to_topics(subject_rows)
            existing = Topic.objects.none()
            if subject is not None:
                existing = Topic.objects.filter(
                    subject=subject, program=Topic.Program.MAARIF,
                ).filter(_grade_filter(grades))
            outcome_count = sum(len(topic['outcomes']) for topic in topics)
            sub_count = sum(
                len(outcome['sub_outcomes'])
                for topic in topics
                for outcome in topic['outcomes']
            )
            plan.append((spec, subject, grades, topics, existing, outcome_count, sub_count))
            bound = 0
            if subject is not None and existing.exists():
                ids = list(existing.values_list('id', flat=True))
                bound = AnswerKeyItem.objects.filter(
                    Q(outcome__topic_id__in=ids) | Q(sub_outcome__outcome__topic_id__in=ids),
                ).count()
            self.stdout.write(
                f"  {code}: {len(topics)} konu, {outcome_count} kazanım, {sub_count} alt kazanım"
                f" — silinecek eski Maarif konu: {existing.count() if subject else 0}"
                f", boşalacak cevap anahtarı bağı: {bound}"
            )

        if options['dry_run']:
            self.stdout.write(self.style.WARNING('Dry-run: veritabanına yazılmadı.'))
            return

        with transaction.atomic():
            for spec, subject, grades, topics, existing, _outcomes, _subs in plan:
                if subject is None:
                    subject = Subject.objects.create(
                        code=spec['code'],
                        name=spec['name'],
                        display_name=spec['display_name'],
                        exam_type_filter=Subject.ExamTypeFilter.ALL,
                        order=spec['order'],
                    )
                kept = {
                    topic.code: topic
                    for topic in existing.filter(id__in=_protected_topic_ids(
                        list(existing.values_list('id', flat=True)),
                    ))
                }
                existing.exclude(id__in=[topic.id for topic in kept.values()]).delete()
                order_base = (
                    Topic.objects.filter(subject=subject).order_by('-order')
                    .values_list('order', flat=True).first() or 0
                )
                for index, topic_data in enumerate(topics, start=1):
                    topic = kept.get(topic_data['code'])
                    if topic is None:
                        topic = Topic.objects.create(
                            subject=subject,
                            program=Topic.Program.MAARIF,
                            code=topic_data['code'],
                            name=topic_data['name'],
                            order=order_base + index,
                        )
                    else:
                        topic.name = topic_data['name']
                        topic.program = Topic.Program.MAARIF
                        topic.order = order_base + index
                        topic.save(update_fields=['name', 'program', 'order'])
                    _write_outcomes(topic, topic_data['outcomes'])
        self.stdout.write(self.style.SUCCESS('Okulizyon Maarif kazanımları yüklendi.'))
